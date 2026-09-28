"""
Lubekalk import export — shared by ``AssignLubekalkPset`` (the CSV artifact) and
``scripts/export_lubekalk_import.py`` (the command line).

Reads each element's ``Lubekalk`` pset, expands it into import rows, sums
identical rows, and writes them the way Lubekalk's *Importera från Bluebeam*
reads them. The limits below were measured on a real import (2026-09-23).
"""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

import ifcopenshell
import ifcopenshell.util.element

from _lubekalk import (
    COLUMNS,
    PSET,
    STATUS_INCLUDED,
    STATUS_OUT_OF_SCOPE,
    STATUS_READY,
    aggregate,
    rows_from_pset,
)

#: Lubekalk's Beteckning. An import row took 30 characters (33+ failed), but a SAK designation is
#: also saved into the article register, which holds 15 ("BDEP-46-040-230" saved, longer ones not).
BETECKNING_MAX = 15
_TAG_ALPHABET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
#: Characters Windows-1252 cannot carry, spelled so the CSV stays readable.
CSV_SPELLING = str.maketrans({"→": ">", "≤": "<=", "≥": ">="})

INSULATION_COLUMNS = ("Isolering", "Kod (kanal/detalj)", "Dim 1", "Dim 2", "P-del", "Montageläge",
                      "Längd", "Antal detaljer")
REVIEW_COLUMNS = ("GlobalId", "Klass", "Namn", "Status", "Säkerhet", "Kod", "Dim 1", "Dim 2",
                  "Orsak", "Källa", "Handling", "P-del")


def _tag(name: str, width: int) -> str:
    digest = int.from_bytes(hashlib.sha1(name.encode("utf-8")).digest()[:8], "big")
    chars = []
    for _ in range(width):
        digest, rest = divmod(digest, len(_TAG_ALPHABET))
        chars.append(_TAG_ALPHABET[rest])
    return "".join(chars)


def short_designations(names: Iterable[str], limit: int = BETECKNING_MAX) -> dict[str, str]:
    """Full → short Beteckning for every name longer than ``limit``; distinct names stay distinct.

    The cut plus ``~`` and a tag computed from the full name, so a product keeps the same short
    name in every export: Lubekalk saves SAK designations into its register, and a name that
    moved between revisions would add a new article each time.
    """
    names = sorted({name for name in names if name})
    taken = {name for name in names if len(name) <= limit}
    short = {}
    for name in names:
        if len(name) <= limit:
            continue
        for width in (2, 3, 4, 6):
            tag = "~" + _tag(name, width)
            candidate = name[:limit - len(tag)].rstrip() + tag
            if candidate not in taken:
                break
        short[name] = candidate
        taken.add(candidate)
    return short


def collect(*files: ifcopenshell.file, omit_codes: Iterable[str] = (),
            beteckning_max: int = BETECKNING_MAX, keep_beteckning: Iterable[str] | None = None) -> dict:
    """One import over one or more patched models (e.g. a discipline's master and its voids).

    Rows whose code is in ``omit_codes`` are moved out of the import into ``omitted``.
    """
    omit = {code.strip() for code in omit_codes if code.strip()}
    rows, review, insulation = [], [], defaultdict(lambda: {"Längd": 0.0, "Antal detaljer": 0})
    status = Counter()
    by_class = defaultdict(Counter)
    codes = Counter()
    for element in (e for file in files for e in file.by_type("IfcElement")):
        pset = ifcopenshell.util.element.get_psets(element).get(PSET)
        if not pset:
            status["(ingen Lubekalk-pset)"] += 1
            by_class[element.is_a()]["(ingen Lubekalk-pset)"] += 1
            continue
        state = pset.get("Status")
        status[state] += 1
        by_class[element.is_a()][state] += 1
        element_rows = rows_from_pset(pset)
        rows.extend(element_rows)
        for row in element_rows:
            codes[row["Kod"]] += 1
        confidence = pset.get("Säkerhet") or ""
        if state not in (STATUS_READY, STATUS_INCLUDED, STATUS_OUT_OF_SCOPE) or (
                state == STATUS_READY and confidence in ("medel", "låg")):
            review.append({
                "GlobalId": element.GlobalId, "Klass": element.is_a(), "Namn": element.Name or "",
                "Status": state, "Säkerhet": confidence, "Kod": pset.get("Kod") or "",
                "Dim 1": pset.get("Dim 1") or "", "Dim 2": pset.get("Dim 2") or "",
                "Orsak": pset.get("Orsak") or "", "Källa": pset.get("Källa") or "",
                "Handling": pset.get("Handling") or "", "P-del": pset.get("P-del") or "",
            })
        if state == STATUS_READY and pset.get("Isolering") and not (
                pset.get("IsoleringKod") or pset.get("DetaljisoleringKod")):
            key = (pset["Isolering"], pset.get("Kod"), pset.get("Dim 1") or 0, pset.get("Dim 2") or 0,
                   pset.get("P-del") or "", pset.get("Montageläge") or "")
            slot = insulation[key]
            slot["Längd"] = round(slot["Längd"] + float(pset.get("Längd") or 0), 3)
            if not pset.get("Längd"):
                slot["Antal detaljer"] += int(pset.get("Antal") or 0)
    insulation_rows = [
        dict(zip(("Isolering", "Kod (kanal/detalj)", "Dim 1", "Dim 2", "P-del", "Montageläge"), key)) | value
        for key, value in sorted(insulation.items(), key=lambda kv: tuple(str(v) for v in kv[0]))
    ]
    total = sum(status.values())
    resolved = status[STATUS_READY] + status[STATUS_INCLUDED] + status[STATUS_OUT_OF_SCOPE]
    if keep_beteckning is not None:
        # Lubekalk hangs a detail under the duct row with the same Beteckning; a note there
        # (round end, insulation type) makes it open an empty duct row instead.
        keep = {code.strip() for code in keep_beteckning if code.strip()}
        for row in rows:
            if row["Kod"] not in keep:
                row["Beteckning"] = ""
    short = short_designations((row.get("Beteckning") or "" for row in rows), beteckning_max)
    for row in rows:
        row["Beteckning"] = short.get(row.get("Beteckning") or "", row.get("Beteckning") or "")
    aggregated = aggregate(rows)
    omitted = [row for row in aggregated if row["Kod"] in omit]
    omitted_totals = defaultdict(lambda: {"rader": 0, "Antal": 0, "Längd": 0.0})
    for row in omitted:
        slot = omitted_totals[row["Kod"]]
        slot["rader"] += 1
        slot["Antal"] += int(row["Antal"] or 0)
        slot["Längd"] = round(slot["Längd"] + float(row["Längd"] or 0), 3)
    return {
        "import": [row for row in aggregated if row["Kod"] not in omit],
        "omitted": omitted,
        "shortened": [{"Beteckning": s, "Fullständig beteckning": full} for full, s in sorted(short.items())],
        "insulation": insulation_rows,
        "review": review,
        "summary": {
            "elements": total,
            "by_status": dict(status),
            "import_ready_share": round(resolved / total, 4) if total else None,
            "rows_before_aggregation": len(rows),
            "codes": dict(codes.most_common()),
            "by_class": {k: dict(v) for k, v in sorted(by_class.items())},
            "review_rows": len(review),
            "insulation_to_code_groups": len(insulation_rows),
            "omitted_codes": dict(sorted(omitted_totals.items())),
            "shortened_designations": len(short),
        },
    }


def _sheet_rows(columns, records):
    return [[record.get(col, "") for col in columns] for record in records]


def _csv_value(value):
    """Swedish number format; whole numbers without decimals (Lubekalk wants 250, not 250,0)."""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).replace(".", ",")
    return value


def write_lubekalk_csv(records: list[dict], path: Path, encoding: str = "cp1252") -> int:
    """The import sheet as Lubekalk reads it. Returns how many characters had to be replaced."""
    lines = [";".join(COLUMNS)]
    for values in _sheet_rows(COLUMNS, records):
        lines.append(";".join(str(_csv_value(v)).replace(";", ",") for v in values))
    text = ("\r\n".join(lines) + "\r\n").translate(CSV_SPELLING)
    data = text.encode(encoding, errors="replace")
    replaced = sum(1 for ch in text if not ch.encode(encoding, errors="ignore"))
    path.write_bytes(data)
    return replaced


def write(result: dict, out: Path) -> list[Path]:
    sheets = {
        "Lubekalk import": (COLUMNS, result["import"]),
        "Isolering att koda": (INSULATION_COLUMNS, result["insulation"]),
        "Granska": (REVIEW_COLUMNS, result["review"]),
    }
    if result.get("shortened"):
        sheets["Förkortade beteckningar"] = (("Beteckning", "Fullständig beteckning"), result["shortened"])
    if result.get("omitted"):
        sheets["Utelämnade koder"] = (COLUMNS, result["omitted"])
    try:
        from openpyxl import Workbook
    except ImportError:
        Workbook = None
    if Workbook is not None and out.suffix.lower() == ".xlsx":
        book = Workbook()
        book.remove(book.active)
        for title, (columns, records) in sheets.items():
            sheet = book.create_sheet(title)
            sheet.append(list(columns))
            for values in _sheet_rows(columns, records):
                sheet.append(values)
            sheet.freeze_panes = "A2"
        status = book.create_sheet("Status")
        for key, value in result["summary"].items():
            status.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value])
        book.save(out)
        return [out]
    written = []
    for title, (columns, records) in sheets.items():
        path = out.with_name(f"{out.stem} - {title}.csv")
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(columns)
            for values in _sheet_rows(columns, records):
                writer.writerow([_csv_value(v) for v in values])
        written.append(path)
    return written


